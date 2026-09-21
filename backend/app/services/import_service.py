"""Servicio de importación de archivos CSV y Excel.

TASK 3316: Implementar la carga de archivos CSV y Excel.

Responsabilidades:
- Detectar el formato del archivo por extensión.
- Leer el contenido binario del archivo.
- Convertir las filas en una lista de diccionarios.
"""

import csv
import io
from typing import Any

from openpyxl import load_workbook


ALLOWED_EXTENSIONS = {".csv", ".xlsx"}


def parse_file(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    """Lee un archivo CSV o Excel y retorna su contenido como lista de diccionarios.

    Cada diccionario representa una fila del archivo, donde las claves
    son los nombres de las columnas del encabezado.

    Args:
        file_bytes: Contenido binario del archivo.
        filename: Nombre original del archivo (se usa para detectar la extensión).

    Returns:
        Lista de diccionarios con los datos de cada fila.

    Raises:
        ValueError: Si la extensión del archivo no es soportada.
    """
    extension = _get_extension(filename)

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Formato de archivo no soportado: '{extension}'. "
            f"Extensiones permitidas: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )

    if extension == ".csv":
        return _parse_csv(file_bytes)

    return _parse_xlsx(file_bytes)


def _get_extension(filename: str) -> str:
    """Extrae la extensión del nombre de archivo en minúsculas.

    Args:
        filename: Nombre original del archivo.

    Returns:
        Extensión con punto incluido, en minúsculas (ej. '.csv').
    """
    dot_index = filename.rfind(".")
    if dot_index == -1:
        return ""
    return filename[dot_index:].lower()


def _parse_csv(file_bytes: bytes) -> list[dict[str, Any]]:
    """Parsea un archivo CSV desde sus bytes.

    Intenta decodificar con utf-8; si falla, utiliza latin-1 como fallback.

    Args:
        file_bytes: Contenido binario del archivo CSV.

    Returns:
        Lista de diccionarios, uno por cada fila de datos.
    """
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1")

    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def _parse_xlsx(file_bytes: bytes) -> list[dict[str, Any]]:
    """Parsea un archivo Excel (.xlsx) desde sus bytes.

    Utiliza la primera fila como encabezados y las siguientes como datos.

    Args:
        file_bytes: Contenido binario del archivo Excel.

    Returns:
        Lista de diccionarios, uno por cada fila de datos.
    """
    workbook = load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if len(rows) < 1:
        return []

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]

    data: list[dict[str, Any]] = []
    for row in rows[1:]:
        row_dict: dict[str, Any] = {}
        for header, cell in zip(headers, row):
            row_dict[header] = cell
        data.append(row_dict)

    return data
